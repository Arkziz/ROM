from __future__ import annotations
import sys,tempfile
from pathlib import Path
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1]
if str(R) not in sys.path:sys.path.insert(0,str(R))
import tools.run_all_class_10_day_diagnostic_matrix as d
events=[{'day':day,'event_type':'victory','details':{'exp':10}} for day in range(1,11)];events.append({'day':10,'event_type':'simulation_final_state','details':{}})
result={'final_level':4,'battles':10,'battle_outcomes':{'victory':10,'defeat':0,'escape':0},'encounters':{'normal':10,'elite':0,'boss':0},'exp_gained':100,'economy':{'repair_cost':0,'final':1},'final_loadout':[],'loot_lifecycle':[],'canonical_result':{'status':'max_days_reached','days':10,'elapsed_minutes':14400,'events':events}}
summary=d.summarize(result,10);assert summary['exp_gained']==100 and 'final_exp' not in summary
data={'days':10,'results':{c:result for c in d.CLASSES},'summaries':{c:summary for c in d.CLASSES},'checks':[]}
with tempfile.TemporaryDirectory() as td:
 p=Path(td)/'d.xlsx';d.write_xlsx(data,p);ws=load_workbook(p)['Diagnostic Summary'];assert ws['D1'].value=='exp_gained' and ws['D2'].value==100
print('[PASS] diagnostic uses available cumulative exp_gained field')
