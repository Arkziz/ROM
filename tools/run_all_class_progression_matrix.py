#!/usr/bin/env python3
from __future__ import annotations
import argparse,json,sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:sys.path.insert(0,str(ROOT))
from systems.simulation.class_progression import ClassProgressionRequest,run_class_progression
CLASSES=('warrior','knight','scout','rogue','mage','priest')
def format_duration(total_minutes):
 total=max(0,int(total_minutes));days,rem=divmod(total,1440);hours,minutes=divmod(rem,60);return f'{days} d {hours} h {minutes} min'
def require(v,msg,checks):
 checks.append({'name':msg,'status':'PASS' if v else 'FAIL'})
 if not v:raise AssertionError(msg)
def validate(data):
 checks=[];results=data['results'];require(tuple(results)==CLASSES,'six class order',checks);signatures=set()
 for cls in CLASSES:
  x=results[cls];out=x['battle_outcomes'];enc=x['encounters'];eco=x['economy']
  require(x['target_reached'] and x['final_level']>=data['target_level'],f'{cls} target reached',checks)
  require(out['victory']+out['defeat']+out['escape']==x['battles'],f'{cls} battle conservation',checks)
  require(enc['normal']+enc['elite']+enc['boss']==x['battles'],f'{cls} encounter conservation',checks)
  require(out['escape']==0,f'{cls} no combat escape',checks)
  require(abs(round(eco['from_mobs']+eco['from_item_sales']-eco['repair_cost'],2)-eco['final'])<.011,f'{cls} economy reconciliation',checks)
  require(all(i['destination']=='equipped' and 0<=i['durability']<=100 for i in x['final_loadout']),f'{cls} final equipment validity',checks)
  for row in x['loot_lifecycle']:require(row['found']==row['equipped']+row['sold']+row['used_or_transformed']+row['remaining'],f"{cls} loot {row['item_id']} conservation",checks)
  signatures.add((x['battles'],out['victory'],out['defeat'],round(eco['final'],2),tuple((i['slot_key'],i['item_id']) for i in x['final_loadout'])))
 require(len(signatures)>1,'class results are not accidental shared state',checks)
 return checks
def summary_rows(data):
 rows=[]
 for cls,x in data['results'].items():
  o=x['battle_outcomes'];e=x['encounters'];eco=x['economy'];res=x['resources'];final_event=next(z for z in reversed(x['canonical_result']['events']) if z['event_type']=='simulation_final_state');fd=final_event['details']
  rows.append([cls,x['final_level'],format_duration(x['time_to_target']['total_minutes']),x['time_to_target']['total_minutes'],x['battles'],o['victory'],o['defeat'],o['escape'],e['normal'],e['elite'],e['boss'],x['exp_gained'],eco['from_mobs'],eco['from_item_sales'],eco['repair_cost'],eco['final'],res['hp_potions']['found'],res['hp_potions']['remaining'],res['refine_stones']['found'],res['refine_stones']['remaining'],res['runes']['found'],res['runes']['remaining'],fd.get('refine_attempts',0),fd.get('refine_successes',0),len(x['final_loadout'])])
 return rows
def style(ws,name):
 ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
 if ws.max_row>=2:
  tab=Table(displayName=name,ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}');tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(tab)
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='1F4E78');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True)
 for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=min(32,max(11,max(len(str(ws.cell(r,col).value or '')) for r in range(1,ws.max_row+1))+2))
def write_xlsx(data,path):
 wb=Workbook();ws=wb.active;ws.title='Matrix Summary';headers=['class','final_level','time_to_level','total_minutes','battles','wins','defeats','escapes','normal','elite','boss','exp','mob_gold','sales_gold','repair_cost','final_gold','potions_found','potions_remaining','stones_found','stones_remaining','runes_found','runes_remaining','refine_attempts','refine_successes','equipped_items'];ws.append(headers)
 for row in summary_rows(data):ws.append(row)
 style(ws,'MatrixSummaryTable')
 eq=wb.create_sheet('Final Equipment');eq.append(['class','simulation_item_id','slot','name','rarity','item_level','refine_level','durability','runes'])
 for cls,x in data['results'].items():
  for i in x['final_loadout']:eq.append([cls,i['simulation_item_id'],i.get('slot_uk') or i['slot_key'],i.get('name_uk') or i['name'],i.get('rarity_uk') or i['rarity'],i['item_level'],i['refine_level'],i['durability'],'; '.join(f"{r['stat']} {r['level']} [slot {r['slot_index']}]" for r in i.get('runes',[]))])
 style(eq,'MatrixEquipmentTable')
 loot=wb.create_sheet('Loot Details');loot.append(['class','item_id','name','category','slot','item_level','rarity','armor_type','stats','found','equipped','sold','used_or_transformed','remaining','sale_gold'])
 for cls,x in data['results'].items():
  for i in x['loot_lifecycle']:loot.append([cls,i['item_id'],i.get('name_uk') or i.get('name_raw'),i.get('category'),i.get('slot_uk') or i.get('slot_key'),i.get('item_level'),i.get('rarity_uk') or i.get('rarity'),i.get('armor_type_uk') or i.get('armor_type'),json.dumps(i.get('stats',{}),ensure_ascii=False,sort_keys=True),i['found'],i['equipped'],i['sold'],i['used_or_transformed'],i['remaining'],i.get('sale_gold',0.0)])
 style(loot,'MatrixLootDetailsTable')
 eqs=wb.create_sheet('Equipment Summary');eqs.append(['class','slot','rarity','item_level','armor_type','found','equipped','sold','used_or_transformed','remaining','sale_gold'])
 for cls,x in data['results'].items():
  for i in x.get('equipment_summary',[]):eqs.append([cls,i.get('slot_uk') or i.get('slot_key'),i.get('rarity_uk') or i.get('rarity'),i.get('item_level'),i.get('armor_type_uk') or i.get('armor_type'),i.get('found',0),i.get('equipped',0),i.get('sold',0),i.get('used_or_transformed',0),i.get('remaining',0),i.get('sale_gold',0.0)])
 style(eqs,'MatrixEquipmentSummaryTable')
 eqd=wb.create_sheet('Equipment Details');eqd.append(['class','item_id','name','slot','rarity','item_level','armor_type','stats','found','equipped','sold','remaining','sale_gold'])
 for cls,x in data['results'].items():
  for i in x.get('equipment_details',[]):eqd.append([cls,i.get('item_id'),i.get('name_uk') or i.get('name_raw'),i.get('slot_uk') or i.get('slot_key'),i.get('rarity_uk') or i.get('rarity'),i.get('item_level'),i.get('armor_type_uk') or i.get('armor_type'),json.dumps(i.get('stats',{}),ensure_ascii=False,sort_keys=True),i.get('found',0),i.get('equipped',0),i.get('sold',0),i.get('remaining',0),i.get('sale_gold',0.0)])
 style(eqd,'MatrixEquipmentDetailsTable')
 val=wb.create_sheet('Validation');val.append(['check','status'])
 for c in data['checks']:val.append([c['name'],c['status']])
 style(val,'MatrixValidationTable');wb.save(path)
def run_matrix(target_level=2,seed=3,max_days=365):
 results={}
 for cls in CLASSES:
  print(f'[RUN] {cls} 1->{target_level} seed {seed}');results[cls]=run_class_progression(ROOT,ClassProgressionRequest(cls,target_level,seed,'hardcore',1,max_days,40));minutes=results[cls]['time_to_target']['total_minutes'];out=results[cls]['battle_outcomes'];print(f"[PASS] {cls} 1->{target_level} | time: {format_duration(minutes)} | battles: {results[cls]['battles']} | wins: {out['victory']} | defeats: {out['defeat']}")
 data={'status':'completed','target_level':target_level,'seed':seed,'classes':list(CLASSES),'results':results};data['checks']=validate(data);data['check_count']=len(data['checks']);data['status']='PASS';print('[TIME RANKING]');
 for rank,(cls,result) in enumerate(sorted(results.items(),key=lambda item:item[1]['time_to_target']['total_minutes']),1):print(f"{rank}. {cls}: {format_duration(result['time_to_target']['total_minutes'])}")
 return data
def main():
 p=argparse.ArgumentParser();p.add_argument('--target-level',type=int,default=2);p.add_argument('--seed',type=int,default=3);p.add_argument('--max-days',type=int,default=365);a=p.parse_args()
 if a.target_level not in (2,5,10):raise SystemExit('Matrix runner supports target levels 2, 5, and 10 only')
 data=run_matrix(a.target_level,a.seed,a.max_days);out=ROOT/'reports'/'simulation';out.mkdir(parents=True,exist_ok=True);stem=f'stage_18_30_all_class_level_1_to_{a.target_level}_seed_{a.seed}';jp=out/(stem+'.json');xp=out/(stem+'.xlsx');jp.write_text(json.dumps(data,ensure_ascii=False,indent=2,sort_keys=True)+'\n',encoding='utf-8');write_xlsx(data,xp)
 print(f"[PASS] all-class matrix checks: {data['check_count']}");print('[REPORT] '+str(jp.relative_to(ROOT)));print('[REPORT] '+str(xp.relative_to(ROOT)));return 0
if __name__=='__main__':raise SystemExit(main())
