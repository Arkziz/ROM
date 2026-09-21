from __future__ import annotations
import sys,tempfile
from pathlib import Path
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1]
if str(R) not in sys.path:sys.path.insert(0,str(R))
import tools.run_all_class_10_day_diagnostic_matrix as d
events=[]
for day in range(1,11):events.extend([{'day':day,'event_type':'victory','details':{'exp':10}},{'day':day,'event_type':'daily_claim','details':{'firius_tokens':5}}])
events.append({'day':10,'event_type':'simulation_final_state','details':{'refine_attempts':0}})
results={}
for cls in d.CLASSES:results[cls]={'final_level':4,'final_exp':123,'battles':10,'battle_outcomes':{'victory':10,'defeat':0,'escape':0},'encounters':{'normal':10,'elite':0,'boss':0},'exp_gained':100,'economy':{'repair_cost':0,'final':10},'final_loadout':[],'loot_lifecycle':[{'item_id':726,'found':50,'equipped':0,'sold':0,'used_or_transformed':0,'remaining':50}],'canonical_result':{'status':'max_days_reached','days':10,'elapsed_minutes':14400,'events':events}}
data={'days':10,'results':results,'summaries':{c:d.summarize(x,10) for c,x in results.items()}};data['checks']=d.validate(data)
with tempfile.TemporaryDirectory() as td:
 p=Path(td)/'d.xlsx';d.write_xlsx(data,p);wb=load_workbook(p);assert wb.sheetnames==['Diagnostic Summary','Daily Diagnostics','Final Equipment','Validation'];assert wb['Daily Diagnostics'].max_row==61
assert data['summaries']['warrior']['daily'][0]['firius_awarded']==5
print('[PASS] fixed 10-day partial-progression diagnostic fixture')
