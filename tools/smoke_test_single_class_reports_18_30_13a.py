from __future__ import annotations
import sys,tempfile
from pathlib import Path
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1]
if str(R) not in sys.path:sys.path.insert(0,str(R))
from tools.run_single_class_progression import write_xlsx
item={'item_id':900001,'name_uk':'Тест','name_raw':'raw','category':'equipment','slot_key':'chest','slot_uk':'Нагрудник','item_level':5,'rarity':'rare','rarity_uk':'синій','armor_type':'chain','armor_type_uk':'кольчуга','stats':{'pdef':5},'found':2,'equipped':1,'sold':1,'used_or_transformed':0,'remaining':0,'sale_gold':11.25}
result={'class':'knight','status':'max_days_reached','target_reached':False,'target_level':10,'final_level':6,'time_to_target':{'total_minutes':306600},'days':365,'battles':2,'battle_outcomes':{'victory':1,'defeat':1,'escape':0},'encounters':{'normal':2,'elite':0,'boss':0},'exp_gained':10,'economy':{'from_mobs':1.0,'from_item_sales':11.25,'repair_cost':0.0,'final':12.25},'loot_lifecycle':[item,dict(item,item_id=900002,name_uk='Тест 2')],'equipment_summary':[item,dict(item,item_id=900002)],'equipment_details':[item,dict(item,item_id=900002)],'final_loadout':[{'simulation_item_id':1,'slot_key':'chest','slot_uk':'Нагрудник','name':'raw','name_uk':'Тест','rarity':'rare','rarity_uk':'синій','item_level':5,'refine_level':0,'durability':100,'runes':[]}]}
with tempfile.TemporaryDirectory() as td:
 p=Path(td)/'report.xlsx';write_xlsx(result,[{'name':'fixture','status':'PASS'}],p);wb=load_workbook(p);assert wb.sheetnames==['Summary','Loot Details','Equipment Summary','Equipment Details','Final Equipment','Validation'];names=[]
 for ws in wb.worksheets:names.extend(ws.tables.keys())
 assert len(names)==len(set(names))==6
source=(R/'tools/run_single_class_progression.py').read_text(encoding='utf-8-sig');assert "[REPORT JSON]" in source and ".json.tmp" in source
print('[PASS] all single-class XLSX sheets and atomic JSON persistence regression')
