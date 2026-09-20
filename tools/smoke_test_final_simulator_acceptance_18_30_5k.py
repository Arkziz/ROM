from __future__ import annotations
import json,sys,tempfile
from pathlib import Path
from openpyxl import Workbook
R=Path(__file__).resolve().parents[1]
if str(R) not in sys.path:sys.path.insert(0,str(R))
from tools.validate_final_simulator_acceptance import validate_pair
fixture={"battle_outcomes":{"victory":1,"defeat":0,"escape":0},"battles":1,"encounters":{"normal":1,"elite":0,"boss":0},"wins":1,"defeats":0,"escapes":0,"target_reached":True,"status":"completed","final_level":2,"target_level":2,"items_awarded":1,"items_lost":0,"loot_lifecycle":[{"item_id":601,"found":1,"equipped":0,"sold":0,"used_or_transformed":1,"remaining":0}],"final_loadout":[],"canonical_result":{"events":[]},"resources":{"hp_potions":{"found":0,"remaining":0}},"economy":{"from_mobs":1.0,"from_item_sales":0.0,"repair_cost":0.0,"final":1.0},"class":"scout","seed":99}
with tempfile.TemporaryDirectory() as d:
 d=Path(d);jp=d/"a.json";xp=d/"a.xlsx";jp.write_text(json.dumps(fixture),encoding="utf-8");wb=Workbook();s=wb.active;s.title="Summary";s.append(["Показник","Значення"])
 for row in [("Клас","scout"),("Seed",99),("Фінальний рівень",2),("Бої",1),("Фінальне золото",1.0),("Золото з мобів",1.0),("Золото з продажу",0.0)]:s.append(row)
 e=wb.create_sheet("Final Equipment");e.append(["simulation_item_id","refine_level","durability","runes"])
 l=wb.create_sheet("Loot Lifecycle");l.append(["item_id"]);l.append([601])
 c=wb.create_sheet("Economy");c.append(["Показник","Сума"])
 for row in [("Стартове золото",0.0),("Золото з мобів",1.0),("Золото з продажу",0.0),("Витрати на ремонт",0.0),("Розрахункове фінальне золото","=SUM(B2:B5)"),("Фактичне фінальне золото",1.0),("Різниця","=B6-B7")]:c.append(row)
 wb.save(xp);result=validate_pair(jp,xp);assert result["status"]=="PASS" and result["check_count"]>=15
print("[PASS] final acceptance gate fixture")
