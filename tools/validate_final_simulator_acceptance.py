#!/usr/bin/env python3
from __future__ import annotations
import argparse,json,math,sys
from pathlib import Path
from openpyxl import load_workbook
REQUIRED_SHEETS=["Summary","Final Equipment","Loot Lifecycle","Economy"]
def close(a,b,tol=0.011):return abs(float(a or 0)-float(b or 0))<=tol
def require(ok,msg,checks):
 checks.append({"name":msg,"status":"PASS" if ok else "FAIL"})
 if not ok:raise AssertionError(msg)
def kv_sheet(ws):return {str(ws.cell(r,1).value):ws.cell(r,2).value for r in range(2,ws.max_row+1) if ws.cell(r,1).value is not None}
def table_rows(ws):
 headers=[c.value for c in ws[1]];return [{headers[i]:row[i].value for i in range(len(headers))} for row in ws.iter_rows(min_row=2) if any(c.value is not None for c in row)]
def validate_pair(json_path:Path,xlsx_path:Path,source_root:Path|None=None):
 data=json.loads(json_path.read_text(encoding="utf-8"));wb=load_workbook(xlsx_path,data_only=False);checks=[]
 require(wb.sheetnames==REQUIRED_SHEETS,"workbook sheet contract",checks)
 battles=int(data["battles"]);out=data["battle_outcomes"];enc=data["encounters"]
 require(int(out["victory"])+int(out["defeat"])+int(out["escape"])==battles,"battle outcome conservation",checks)
 require(sum(int(enc[k]) for k in ("normal","elite","boss"))==battles,"encounter conservation",checks)
 require(int(data["wins"])==int(out["victory"]) and int(data["defeats"])==int(out["defeat"]) and int(data["escapes"])==int(out["escape"]),"top-level battle parity",checks)
 require(bool(data["target_reached"]) and data["status"]=="completed" and int(data["final_level"])>=int(data["target_level"]),"target completion",checks)
 require(int(data["items_awarded"])+int(data["items_lost"])>=sum(int(x["found"]) for x in data["loot_lifecycle"]),"loot accounting coverage",checks)
 for row in data["loot_lifecycle"]:
  f=int(row["found"]);parts=int(row["equipped"])+int(row["sold"])+int(row["used_or_transformed"])+int(row["remaining"])
  require(parts==f,f"loot conservation item {row['item_id']}",checks)
 final={int(x["simulation_item_id"]):x for x in data["final_loadout"]}
 require(len(final)==len(data["final_loadout"]),"unique final equipment instances",checks)
 for item in final.values():
  require(item["destination"]=="equipped" and int(item["durability"])>=0,"final equipment state",checks)
  require(len(item.get("runes",[]))==len({int(r["slot_index"]) for r in item.get("runes",[])}),"unique rune sockets",checks)
 events=data["canonical_result"]["events"]
 deaths=[e for e in events if e["event_type"]=="death"]
 require(len(deaths)==int(data["defeats"]),"death event parity",checks)
 require(all(int((e.get("details") or {}).get("durability_items",0))>=0 for e in deaths),"durability event accounting",checks)
 pots=sum(1 for e in events if e["event_type"]=="potion_used");found=int((data["resources"].get("hp_potions") or {}).get("found",0));remaining=int((data["resources"].get("hp_potions") or {}).get("remaining",0))
 require(found-remaining==pots,"potion conservation",checks)
 waits=[e for e in events if e["event_type"]=="pre_hunt_hp_wait"]
 require(all(int(e["details"]["threshold_percent"])==25 for e in waits),"pre-hunt HP wait threshold",checks)
 require(all(int(e["details"]["threshold_percent"])==40 for e in events if e["event_type"]=="potion_used"),"auto-potion threshold",checks)
 economy=data["economy"];expected=round(float(economy["from_mobs"])+float(economy["from_item_sales"])-float(economy["repair_cost"]),2)
 require(close(expected,economy["final"]),"economy reconciliation",checks)
 summary=kv_sheet(wb["Summary"])
 require(summary.get("Клас")==data["class"] and int(summary.get("Seed"))==int(data["seed"]),"summary identity parity",checks)
 require(int(summary.get("Фінальний рівень"))==int(data["final_level"]) and int(summary.get("Бої"))==battles,"summary progression parity",checks)
 require(close(summary.get("Фінальне золото"),economy["final"]) and close(summary.get("Золото з мобів"),economy["from_mobs"]) and close(summary.get("Золото з продажу"),economy["from_item_sales"]),"summary economy parity",checks)
 eq_rows=table_rows(wb["Final Equipment"]);eq={int(x["simulation_item_id"]):x for x in eq_rows}
 require(set(eq)==set(final),"final equipment instance parity",checks)
 for key,item in final.items():
  row=eq[key];require(int(row["refine_level"])==int(item["refine_level"]) and int(row["durability"])==int(item["durability"]),f"equipment refine/durability {key}",checks)
  rune_text=str(row.get("runes") or "");require(all(f"{r['stat']} {r['level']}" in rune_text and f"slot {r['slot_index']}" in rune_text for r in item.get("runes",[])),f"equipment rune parity {key}",checks)
 loot_rows=table_rows(wb["Loot Lifecycle"]);loot={int(x["item_id"]):x for x in loot_rows}
 require(set(loot)=={int(x["item_id"]) for x in data["loot_lifecycle"]},"loot workbook item parity",checks)
 eco=kv_sheet(wb["Economy"]);require(close(eco["Золото з мобів"],economy["from_mobs"]) and close(eco["Золото з продажу"],economy["from_item_sales"]) and close(eco["Фактичне фінальне золото"],economy["final"]),"economy workbook parity",checks)
 require(eco["Розрахункове фінальне золото"]=="=SUM(B2:B5)" and eco["Різниця"]=="=B6-B7","economy formula contract",checks)
 if source_root:
  runner=(source_root/"tools/run_class_progression_simulation.py").read_text(encoding="utf-8-sig")
  require("Звіти збережено (2 файли)" in runner and "_loot.csv" not in runner and "_equipment.json" not in runner,"two-output source contract",checks)
  policy=(source_root/"systems/simulation/behavior_policy.py").read_text(encoding="utf-8-sig")
  survival=(source_root/"systems/simulation/survival_state.py").read_text(encoding="utf-8-sig")
  require("repair_threshold: int = 15" in policy and "min(values) <= self.repair_threshold" in policy,"repair maintenance boundary",checks)
  require("item.durability<=int(threshold)" in survival and "old-5" in survival,"repair/death durability contract",checks)
 return {"status":"PASS","json":str(json_path),"xlsx":str(xlsx_path),"checks":checks,"check_count":len(checks)}
def main():
 p=argparse.ArgumentParser();p.add_argument("--json",required=True);p.add_argument("--xlsx",required=True);p.add_argument("--report");p.add_argument("--source-root");a=p.parse_args();result=validate_pair(Path(a.json),Path(a.xlsx),Path(a.source_root) if a.source_root else None)
 if a.report:Path(a.report).write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 print(f"[PASS] final simulator acceptance checks: {result['check_count']}");return 0
if __name__=="__main__":raise SystemExit(main())
