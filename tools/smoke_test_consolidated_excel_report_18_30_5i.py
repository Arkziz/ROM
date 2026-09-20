from __future__ import annotations
import json,sys,tempfile
from pathlib import Path
R=Path(__file__).resolve().parents[1]
if str(R) not in sys.path:sys.path.insert(0,str(R))
from openpyxl import load_workbook
from systems.simulation.excel_reporting import write_consolidated_workbook
fixture={"class":"scout","seed":2,"start_level":1,"final_level":2,"target_reached":True,"time_to_target":{"formatted":"0 дн 8 год 55 хв"},"battles":49,"wins":41,"defeats":8,"escapes":0,"encounters":{"normal":46,"elite":3,"boss":0},"exp_gained":504,"economy":{"final":4.98,"from_mobs":4.68,"from_item_sales":0.30,"repair_cost":0.0},"resources":{"refine_stones":{"found":2,"remaining":0},"runes":{"found":2,"remaining":0}},"canonical_result":{"events":[{"event_type":"resource_maintenance","details":{"runes_inserted":1}},{"event_type":"simulation_final_state","details":{"refine_attempts":2,"refine_successes":2}}]},"final_loadout":[{"simulation_item_id":5,"item_id":900005,"slot_uk":"Нагрудник","name_uk":"Нагрудник +2","rarity_uk":"синій","item_level":1,"refine_level":2,"durability":75,"stats":{"sta":1,"pdef":1,"mdef":1},"runes":[{"slot_index":0,"rune_item_id":711,"stat":"DEX","level":1}]}],"loot_lifecycle":[{"item_id":601,"name_uk":"Камінь Посилення","category":"material","found":2,"equipped":0,"sold":0,"used_or_transformed":2,"remaining":0,"sale_gold":0.0}]}
with tempfile.TemporaryDirectory() as d:
 p=write_consolidated_workbook(fixture,Path(d)/"report.xlsx");wb=load_workbook(p,data_only=False);assert wb.sheetnames==["Summary","Final Equipment","Loot Lifecycle","Economy"];eq=wb["Final Equipment"];assert eq["G2"].value==2 and eq["H2"].value==75 and "DEX 1" in eq["R2"].value;eco=wb["Economy"];assert eco["B6"].value=="=SUM(B2:B5)" and eco["B7"].value==4.98;assert p.stat().st_size>5000
print("[PASS] consolidated workbook sheets, final refine/rune/durability and economy reconciliation")
