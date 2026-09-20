from __future__ import annotations
import sys,tempfile
from pathlib import Path
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1]
if str(R) not in sys.path:sys.path.insert(0,str(R))
from tools.run_all_class_progression_matrix import CLASSES,validate,write_xlsx
results={}
for n,cls in enumerate(CLASSES):
 results[cls]={'target_reached':True,'final_level':2,'battles':2+n,'battle_outcomes':{'victory':2+n,'defeat':0,'escape':0},'encounters':{'normal':2+n,'elite':0,'boss':0},'economy':{'from_mobs':1+n,'from_item_sales':.3,'repair_cost':0,'final':1.3+n},'final_loadout':[],'loot_lifecycle':[{'item_id':100+n,'found':1,'equipped':0,'sold':0,'used_or_transformed':0,'remaining':1}],'time_to_target':{'total_minutes':10+n},'exp_gained':10,'resources':{'hp_potions':{'found':0,'remaining':0},'refine_stones':{'found':0,'remaining':0},'runes':{'found':0,'remaining':0}},'canonical_result':{'events':[{'event_type':'simulation_final_state','details':{'refine_attempts':0,'refine_successes':0}}]}}
data={'target_level':2,'seed':3,'results':results};data['checks']=validate(data)
with tempfile.TemporaryDirectory() as d:
 p=Path(d)/'matrix.xlsx';write_xlsx(data,p);wb=load_workbook(p);assert wb.sheetnames==['Matrix Summary','Final Equipment','Resources','Validation'];assert wb['Matrix Summary'].max_row==7
print('[PASS] six-class matrix aggregate JSON/XLSX fixture')
