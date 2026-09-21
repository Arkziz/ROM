#!/usr/bin/env python3
from __future__ import annotations
import argparse,json,sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:sys.path.insert(0,str(ROOT))
from systems.simulation.class_progression import ClassProgressionRequest,run_class_progression
CLASSES=('warrior','knight','scout','rogue','mage','priest');DIAGNOSTIC_TARGET=50
def duration(total):
 d,r=divmod(max(0,int(total)),1440);h,m=divmod(r,60);return f'{d} d {h} h {m} min'
def require(value,name,checks):
 checks.append({'name':name,'status':'PASS' if value else 'FAIL'})
 if not value:raise AssertionError(name)
def build_daily(result,days):
 rows=[];events=result['canonical_result']['events']
 for day in range(1,days+1):
  ev=[e for e in events if int(e.get('day',0))==day];wins=sum(e['event_type']=='victory' for e in ev);losses=sum(e['event_type']=='death' for e in ev);levelups=[int(e['details']['level']) for e in ev if e['event_type']=='level_up'];repairs=[e for e in ev if e['event_type']=='repair'];daily=[e for e in ev if e['event_type']=='daily_claim'];broken=max([int(e.get('details',{}).get('broken',0)) for e in ev if e['event_type']=='death'] or [0])
  rows.append({'day':day,'wins':wins,'defeats':losses,'battles':wins+losses,'win_rate':round(100*wins/(wins+losses),2) if wins+losses else 0.0,'exp_won':sum(int(e.get('details',{}).get('exp',0)) for e in ev if e['event_type']=='victory'),'level_ups':levelups,'daily_claims':len(daily),'firius_awarded':sum(int(e.get('details',{}).get('firius_tokens',0)) for e in daily),'repair_events':len(repairs),'repair_cost':round(sum(float(e.get('details',{}).get('cost',0)) for e in repairs),2),'max_broken_items':broken})
 return rows
def summarize(result,days):
 out=result['battle_outcomes'];events=result['canonical_result']['events'];final=next(e['details'] for e in reversed(events) if e['event_type']=='simulation_final_state');daily=build_daily(result,days);durabilities=[int(i.get('durability',0)) for i in result['final_loadout']]
 return {'final_level':result['final_level'],'exp_gained':result['exp_gained'],'battles':result['battles'],'wins':out['victory'],'defeats':out['defeat'],'win_rate':round(100*out['victory']/result['battles'],2) if result['battles'] else 0.0,'exp_gained':result['exp_gained'],'exp_per_battle':round(result['exp_gained']/result['battles'],3) if result['battles'] else 0.0,'exp_per_day':round(result['exp_gained']/days,2),'elapsed_minutes':result['canonical_result']['elapsed_minutes'],'elapsed':duration(result['canonical_result']['elapsed_minutes']),'repair_cost':result['economy']['repair_cost'],'final_gold':result['economy']['final'],'equipped_items':len(result['final_loadout']),'broken_items':sum(x==0 for x in durabilities),'low_durability_items':sum(x<=15 for x in durabilities),'daily':daily,'final_state':final}
def validate(data):
 checks=[];require(tuple(data['results'])==CLASSES,'six class order',checks)
 for cls,result in data['results'].items():
  out=result['battle_outcomes'];enc=result['encounters'];require(result['canonical_result']['status']=='max_days_reached',f'{cls} stopped by day cap',checks);require(result['canonical_result']['days']==data['days'],f'{cls} exact day count',checks);require(out['victory']+out['defeat']+out['escape']==result['battles'],f'{cls} battle conservation',checks);require(enc['normal']+enc['elite']+enc['boss']==result['battles'],f'{cls} encounter conservation',checks);require(out['escape']==0,f'{cls} no combat escape',checks);require(len(data['summaries'][cls]['daily'])==data['days'],f'{cls} daily rows',checks)
  for row in result['loot_lifecycle']:require(row['found']==row['equipped']+row['sold']+row['used_or_transformed']+row['remaining'],f"{cls} loot {row['item_id']} conservation",checks)
 return checks
def style(ws,name):
 ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
 if ws.max_row>=2:
  tab=Table(displayName=name,ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}');tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(tab)
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='1F4E78');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True)
 for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=min(34,max(11,max(len(str(ws.cell(r,col).value or '')) for r in range(1,ws.max_row+1))+2))
def write_xlsx(data,path):
 wb=Workbook();ws=wb.active;ws.title='Diagnostic Summary';headers=['class','days','final_level','exp_gained','elapsed','battles','wins','defeats','win_rate','exp_gained','exp_per_battle','exp_per_day','repair_cost','final_gold','equipped_items','broken_items','low_durability_items'];ws.append(headers)
 for cls,s in data['summaries'].items():ws.append([cls,data['days'],s['final_level'],s['exp_gained'],s['elapsed'],s['battles'],s['wins'],s['defeats'],s['win_rate'],s['exp_gained'],s['exp_per_battle'],s['exp_per_day'],s['repair_cost'],s['final_gold'],s['equipped_items'],s['broken_items'],s['low_durability_items']])
 style(ws,'DiagnosticSummaryTable');daily=wb.create_sheet('Daily Diagnostics');daily.append(['class','day','battles','wins','defeats','win_rate','exp_won','level_ups','daily_claims','firius_awarded','repair_events','repair_cost','max_broken_items'])
 for cls,s in data['summaries'].items():
  for x in s['daily']:daily.append([cls,x['day'],x['battles'],x['wins'],x['defeats'],x['win_rate'],x['exp_won'],','.join(map(str,x['level_ups'])),x['daily_claims'],x['firius_awarded'],x['repair_events'],x['repair_cost'],x['max_broken_items']])
 style(daily,'DailyDiagnosticsTable');eq=wb.create_sheet('Final Equipment');eq.append(['class','slot','name','level','rarity','refine','durability','runes'])
 for cls,result in data['results'].items():
  for i in result['final_loadout']:eq.append([cls,i.get('slot_uk') or i['slot_key'],i.get('name_uk') or i['name'],i['item_level'],i.get('rarity_uk') or i['rarity'],i['refine_level'],i['durability'],'; '.join(f"{r['stat']} {r['level']}" for r in i.get('runes',[]))])
 style(eq,'DiagnosticEquipmentTable');val=wb.create_sheet('Validation');val.append(['check','status'])
 for x in data['checks']:val.append([x['name'],x['status']])
 style(val,'DiagnosticValidationTable');wb.save(path)
def run_diagnostic(days=10,seed=3):
 results={};summaries={}
 for cls in CLASSES:
  print(f'[RUN] {cls} 10-day seed {seed}');result=run_class_progression(ROOT,ClassProgressionRequest(cls,DIAGNOSTIC_TARGET,seed,'hardcore',1,days,40));results[cls]=result;summaries[cls]=summarize(result,days);s=summaries[cls];print(f"[PASS] {cls} | level: {s['final_level']} | exp: {s['exp_gained']} | wins: {s['wins']} | defeats: {s['defeats']} | win rate: {s['win_rate']}% | broken: {s['broken_items']}")
 data={'status':'completed','mode':'fixed_days','days':days,'seed':seed,'diagnostic_target':DIAGNOSTIC_TARGET,'results':results,'summaries':summaries};data['checks']=validate(data);data['check_count']=len(data['checks']);data['status']='PASS';return data
def main():
 p=argparse.ArgumentParser();p.add_argument('--days',type=int,default=10);p.add_argument('--seed',type=int,default=3);a=p.parse_args()
 if a.days!=10:raise SystemExit('Stage 18.30.9 diagnostic runner is fixed to 10 days')
 data=run_diagnostic(a.days,a.seed);out=ROOT/'reports'/'simulation';out.mkdir(parents=True,exist_ok=True);stem=f'stage_18_30_all_class_{a.days}_day_seed_{a.seed}';jp=out/(stem+'.json');xp=out/(stem+'.xlsx');jp.write_text(json.dumps(data,ensure_ascii=False,indent=2,sort_keys=True)+'\n',encoding='utf-8');write_xlsx(data,xp);print(f"[PASS] all-class 10-day diagnostic checks: {data['check_count']}");print('[REPORT] '+str(jp.relative_to(ROOT)));print('[REPORT] '+str(xp.relative_to(ROOT)));return 0
if __name__=='__main__':raise SystemExit(main())
