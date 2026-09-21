#!/usr/bin/env python3
from __future__ import annotations
import argparse,json,sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:sys.path.insert(0,str(ROOT))
from systems.simulation.class_progression import ClassProgressionRequest,run_class_progression
CLASSES=('warrior','knight','scout','rogue','mage','priest')
def duration(minutes):
 d,r=divmod(max(0,int(minutes)),1440);h,m=divmod(r,60);return f'{d} d {h} h {m} min'
def style(ws,name):
 ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
 if ws.max_row>=2:
  t=Table(displayName=name,ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}');t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='1F4E78');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True)
 for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=min(38,max(11,max(len(str(ws.cell(r,col).value or '')) for r in range(1,ws.max_row+1))+2))
def validate(x,target):
 checks=[]
 def req(v,n):checks.append({'name':n,'status':'PASS' if v else 'FAIL'});assert v,n
 o=x['battle_outcomes'];e=x['encounters'];eco=x['economy'];req(o['victory']+o['defeat']+o['escape']==x['battles'],'battle conservation');req(e['normal']+e['elite']+e['boss']==x['battles'],'encounter conservation');req(o['escape']==0,'no combat escape');req(abs(round(eco['from_mobs']+eco['from_item_sales']-eco['repair_cost'],2)-eco['final'])<.011,'economy reconciliation')
 for row in x['loot_lifecycle']:req(row['found']==row['equipped']+row['sold']+row['used_or_transformed']+row['remaining'],f"loot {row['item_id']} conservation")
 if x['target_reached']:req(x['final_level']==target,'exact target reached')
 else:req(x['status']=='max_days_reached','valid max-days partial result')
 req('equipment_details' in x and 'equipment_summary' in x,'expanded equipment export');return checks
def write_xlsx(result,checks,path):
 wb=Workbook();ws=wb.active;ws.title='Summary';o=result['battle_outcomes'];e=result['encounters'];eco=result['economy'];ws.append(['class','run_status','target_reached','target_level','final_level','time','total_minutes','days','battles','wins','defeats','win_rate','normal','elite','boss','exp_gained','mob_gold','sales_gold','repair_cost','final_gold']);mins=result['time_to_target']['total_minutes'];ws.append([result['class'],result['status'],result['target_reached'],result['target_level'],result['final_level'],duration(mins),mins,result['days'],result['battles'],o['victory'],o['defeat'],round(100*o['victory']/result['battles'],2) if result['battles'] else 0,e['normal'],e['elite'],e['boss'],result['exp_gained'],eco['from_mobs'],eco['from_item_sales'],eco['repair_cost'],eco['final']]);style(ws,'SingleClassSummary')
 loot=wb.create_sheet('Loot Details');loot.append(['item_id','name','category','slot','item_level','rarity','armor_type','stats','found','equipped','sold','used_or_transformed','remaining','sale_gold'])
 for i in result['loot_lifecycle']:loot.append([i['item_id'],i.get('name_uk') or i.get('name_raw'),i.get('category'),i.get('slot_uk') or i.get('slot_key'),i.get('item_level'),i.get('rarity_uk') or i.get('rarity'),i.get('armor_type_uk') or i.get('armor_type'),json.dumps(i.get('stats',{}),ensure_ascii=False,sort_keys=True),i['found'],i['equipped'],i['sold'],i['used_or_transformed'],i['remaining'],i.get('sale_gold',0)]);style(loot,'SingleClassLoot')
 eqs=wb.create_sheet('Equipment Summary');eqs.append(['slot','rarity','item_level','armor_type','found','equipped','sold','used_or_transformed','remaining','sale_gold'])
 for i in result.get('equipment_summary',[]):eqs.append([i.get('slot_uk') or i.get('slot_key'),i.get('rarity_uk') or i.get('rarity'),i.get('item_level'),i.get('armor_type_uk') or i.get('armor_type'),i.get('found',0),i.get('equipped',0),i.get('sold',0),i.get('used_or_transformed',0),i.get('remaining',0),i.get('sale_gold',0)]);style(eqs,'SingleClassEquipmentSummary')
 eqd=wb.create_sheet('Equipment Details');eqd.append(['item_id','name','slot','rarity','item_level','armor_type','stats','found','equipped','sold','remaining','sale_gold'])
 for i in result.get('equipment_details',[]):eqd.append([i.get('item_id'),i.get('name_uk') or i.get('name_raw'),i.get('slot_uk') or i.get('slot_key'),i.get('rarity_uk') or i.get('rarity'),i.get('item_level'),i.get('armor_type_uk') or i.get('armor_type'),json.dumps(i.get('stats',{}),ensure_ascii=False,sort_keys=True),i.get('found',0),i.get('equipped',0),i.get('sold',0),i.get('remaining',0),i.get('sale_gold',0)]);style(eqd,'SingleClassEquipmentDetails')
 final=wb.create_sheet('Final Equipment');final.append(['simulation_item_id','slot','name','rarity','item_level','refine_level','durability','runes'])
 for i in result['final_loadout']:final.append([i['simulation_item_id'],i.get('slot_uk') or i['slot_key'],i.get('name_uk') or i['name'],i.get('rarity_uk') or i['rarity'],i['item_level'],i['refine_level'],i['durability'],'; '.join(f"{r['stat']} {r['level']}" for r in i.get('runes',[]))]);style(final,'SingleClassFinalEquipment')
 val=wb.create_sheet('Validation');val.append(['check','status'])
 for c in checks:val.append([c['name'],c['status']])
 style(val,'SingleClassValidation');wb.save(path)
def main():
 p=argparse.ArgumentParser();p.add_argument('--class',dest='character_class',required=True,choices=CLASSES);p.add_argument('--target-level',type=int,required=True,choices=range(2,11),metavar='2..10');p.add_argument('--seed',type=int,default=3);p.add_argument('--max-days',type=int,default=365);a=p.parse_args();print(f'[RUN] {a.character_class} 1->{a.target_level} seed {a.seed}')
 x=run_class_progression(ROOT,ClassProgressionRequest(a.character_class,a.target_level,a.seed,'hardcore',1,a.max_days,40));checks=validate(x,a.target_level);status='PASS' if x['target_reached'] else 'CAP';o=x['battle_outcomes'];print(f"[{status}] {a.character_class} 1->{a.target_level} | final level: {x['final_level']} | time: {duration(x['time_to_target']['total_minutes'])} | battles: {x['battles']} | wins: {o['victory']} | defeats: {o['defeat']}")
 out=ROOT/'reports'/'simulation'/'single_class';out.mkdir(parents=True,exist_ok=True);stem=f"stage_18_30_{a.character_class}_level_1_to_{a.target_level}_seed_{a.seed}";data={'status':'PASS' if x['target_reached'] else 'PARTIAL','class':a.character_class,'target_level':a.target_level,'seed':a.seed,'max_days':a.max_days,'result':x,'checks':checks,'check_count':len(checks)};jp=out/(stem+'.json');xp=out/(stem+'.xlsx');jp.write_text(json.dumps(data,ensure_ascii=False,indent=2,sort_keys=True)+'\n',encoding='utf-8');write_xlsx(x,checks,xp);print('[REPORT] '+str(jp.relative_to(ROOT)));print('[REPORT] '+str(xp.relative_to(ROOT)));return 0
if __name__=='__main__':raise SystemExit(main())
