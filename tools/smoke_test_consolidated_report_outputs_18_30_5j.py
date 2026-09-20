from __future__ import annotations
import ast,subprocess,sys,tempfile
from pathlib import Path
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1]
if str(R) not in sys.path:sys.path.insert(0,str(R))
from systems.simulation.excel_reporting import write_consolidated_workbook
source=(R/"tools"/"run_class_progression_simulation.py").read_text(encoding="utf-8-sig");tree=ast.parse(source);compile(source,"runner","exec")
assert "Звіти збережено (2 файли)" in source
for forbidden in ("_loot.csv","_equipment.json","cp.open(","ep.write_text(","cp.relative_to(ROOT)","ep.relative_to(ROOT)"):assert forbidden not in source
assert source.count("write_consolidated_workbook(data,xp)")==1 and source.count("jp.relative_to(ROOT)")==1 and source.count("xp.relative_to(ROOT)")==1
fixture={"class":"scout","seed":2,"start_level":1,"final_level":2,"target_reached":True,"time_to_target":{"formatted":"0 дн 8 год 55 хв"},"battles":49,"wins":41,"defeats":8,"escapes":0,"encounters":{"normal":46,"elite":3,"boss":0},"exp_gained":504,"economy":{"final":4.98,"from_mobs":4.68,"from_item_sales":0.30,"repair_cost":0.0},"resources":{"refine_stones":{"found":2,"remaining":0},"runes":{"found":2,"remaining":0}},"canonical_result":{"events":[{"event_type":"resource_maintenance","details":{"runes_inserted":1}},{"event_type":"simulation_final_state","details":{"refine_attempts":2,"refine_successes":2}}]},"final_loadout":[{"simulation_item_id":5,"item_id":900005,"slot_uk":"Нагрудник","name_uk":"Нагрудник +2","rarity_uk":"синій","item_level":1,"refine_level":2,"durability":75,"stats":{"sta":1,"pdef":1,"mdef":1},"runes":[{"slot_index":0,"rune_item_id":711,"stat":"DEX","level":1}]}],"loot_lifecycle":[{"item_id":601,"name_uk":"Камінь Посилення","category":"material","found":2,"equipped":0,"sold":0,"used_or_transformed":2,"remaining":0,"sale_gold":0.0}]}
with tempfile.TemporaryDirectory() as d:
 d=Path(d);jp=d/"simulation.json";xp=d/"simulation.xlsx";jp.write_text(__import__("json").dumps(fixture,ensure_ascii=False),encoding="utf-8");write_consolidated_workbook(fixture,xp)
 assert sorted(p.suffix for p in d.iterdir())==[".json",".xlsx"]
 wb=load_workbook(xp,data_only=False);assert wb.sheetnames==["Summary","Final Equipment","Loot Lifecycle","Economy"];assert wb["Final Equipment"]["G2"].value==2 and "DEX 1" in wb["Final Equipment"]["R2"].value;assert wb["Economy"]["B7"].value==4.98
print("[PASS] exactly two report outputs: authoritative JSON and consolidated XLSX")
