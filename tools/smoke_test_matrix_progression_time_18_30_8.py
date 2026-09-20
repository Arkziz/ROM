from __future__ import annotations
import contextlib,io,sys,tempfile
from pathlib import Path
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1]
if str(R) not in sys.path:sys.path.insert(0,str(R))
import tools.run_all_class_progression_matrix as matrix
assert matrix.format_duration(278)=='0 d 4 h 38 min'
assert matrix.format_duration(846)=='0 d 14 h 6 min'
results={}
for n,cls in enumerate(matrix.CLASSES):
 minutes=(6-n)*60+n
 results[cls]={'target_reached':True,'final_level':2,'battles':2+n,'battle_outcomes':{'victory':2+n,'defeat':0,'escape':0},'encounters':{'normal':2+n,'elite':0,'boss':0},'economy':{'from_mobs':1+n,'from_item_sales':.3,'repair_cost':0,'final':1.3+n},'final_loadout':[],'loot_lifecycle':[{'item_id':100+n,'found':1,'equipped':0,'sold':0,'used_or_transformed':0,'remaining':1,'sale_gold':0}],'time_to_target':{'total_minutes':minutes},'exp_gained':10,'resources':{'hp_potions':{'found':0,'remaining':0},'refine_stones':{'found':0,'remaining':0},'runes':{'found':0,'remaining':0}},'canonical_result':{'events':[{'event_type':'simulation_final_state','details':{'refine_attempts':0,'refine_successes':0}}]}}
data={'target_level':2,'seed':3,'results':results};data['checks']=matrix.validate(data)
with tempfile.TemporaryDirectory() as d:
 p=Path(d)/'matrix.xlsx';matrix.write_xlsx(data,p);wb=load_workbook(p);ws=wb['Matrix Summary'];headers=[c.value for c in ws[1]];assert headers[:4]==['class','final_level','time_to_level','total_minutes'];assert ws['C2'].value=='0 d 6 h 0 min';assert ws['D2'].value==360
original=matrix.run_class_progression
try:
 matrix.run_class_progression=lambda root,request:results[request.character_class]
 output=io.StringIO()
 with contextlib.redirect_stdout(output):matrix.run_matrix(2,3,365)
 text=output.getvalue();assert '| time: 0 d 6 h 0 min | battles: 2 | wins: 2 | defeats: 0' in text;assert '[TIME RANKING]' in text;assert '1. priest: 0 d 1 h 5 min' in text
finally:matrix.run_class_progression=original
print('[PASS] console duration, ranking, and XLSX time columns')
