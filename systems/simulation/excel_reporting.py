from __future__ import annotations
from pathlib import Path
from typing import Any
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table,TableStyleInfo
BLUE="1F4E78";LIGHT="D9EAF7";GREEN="008000";GRAY="666666";WHITE="FFFFFF";ORANGE="F4B183"
def _value(v):
 if isinstance(v,(dict,list,tuple)):return "; ".join(f"{k}={x}" for k,x in v.items()) if isinstance(v,dict) else "; ".join(map(str,v))
 return v
def _setup(ws,title):
 ws.sheet_view.showGridLines=False;ws.freeze_panes="A2";ws.auto_filter.ref=None
 ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.oddHeader.center.text=title

def _table(ws,name):
 if ws.max_row<2 or ws.max_column<1:return
 ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}";tab=Table(displayName=name,ref=ref);tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showFirstColumn=False,showLastColumn=False);ws.add_table(tab)
 for cell in ws[1]:cell.fill=PatternFill("solid",fgColor=BLUE);cell.font=Font(color=WHITE,bold=True);cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
 ws.row_dimensions[1].height=32

def _widths(ws):
 for col in range(1,ws.max_column+1):
  values=[str(ws.cell(r,col).value or "") for r in range(1,min(ws.max_row,200)+1)];ws.column_dimensions[get_column_letter(col)].width=min(45,max(10,max(map(len,values),default=10)+2))
 for row in ws.iter_rows():
  for c in row:c.alignment=Alignment(vertical="top",wrap_text=True)

def _runes(item):return "; ".join(f"{r.get('stat','')} {r.get('level','')} [slot {r.get('slot_index','')}]" for r in item.get("runes",[]) or [])
def write_consolidated_workbook(payload:dict[str,Any],output_path:Path|str)->Path:
 output=Path(output_path);output.parent.mkdir(parents=True,exist_ok=True);wb=Workbook();ws=wb.active;ws.title="Summary";_setup(ws,"ROM Simulation Summary")
 economy=dict(payload.get("economy") or {});resources=dict(payload.get("resources") or {});canonical=dict(payload.get("canonical_result") or {})
 summary=[("Клас",payload.get("class")),("Seed",payload.get("seed")),("Початковий рівень",payload.get("start_level")),("Фінальний рівень",payload.get("final_level")),("Ціль досягнута",payload.get("target_reached")),("Фактичний час",(payload.get("time_to_target") or {}).get("formatted")),("Бої",payload.get("battles")),("Перемоги",payload.get("wins")),("Поразки",payload.get("defeats")),("Втечі",payload.get("escapes")),("Normal",(payload.get("encounters") or {}).get("normal")),("Elite",(payload.get("encounters") or {}).get("elite")),("Boss",(payload.get("encounters") or {}).get("boss")),("EXP",payload.get("exp_gained")),("Фінальне золото",economy.get("final")),("Золото з мобів",economy.get("from_mobs")),("Золото з продажу",economy.get("from_item_sales")),("Витрати на ремонт",economy.get("repair_cost")),("Спроби заточки",_final_detail(canonical,"refine_attempts")),("Успіхи заточки",_final_detail(canonical,"refine_successes")),("Вставлено рун",sum(int((e.get("details") or {}).get("runes_inserted",0)) for e in canonical.get("events",[]) if e.get("event_type") in {"resource_maintenance","city_visit"})),("Камені знайдено",(resources.get("refine_stones") or {}).get("found")),("Камені залишились",(resources.get("refine_stones") or {}).get("remaining")),("Руни знайдено",(resources.get("runes") or {}).get("found")),("Руни залишились",(resources.get("runes") or {}).get("remaining"))]
 ws.append(["Показник","Значення"])
 for row in summary:ws.append(row)
 _table(ws,"SummaryTable");_widths(ws);ws.column_dimensions["A"].width=28;ws.column_dimensions["B"].width=24
 eq=wb.create_sheet("Final Equipment");_setup(eq,"Final Equipment");headers=["simulation_item_id","item_id","slot","name","rarity","item_level","refine_level","durability","STR","DEX","INT","WIS","STA","PDEF","MDEF","weapon_min","weapon_max","runes","rune_slots_used"] ;eq.append(headers)
 for item in payload.get("final_loadout",[]) or []:
  st=item.get("stats") or {};eq.append([item.get("simulation_item_id"),item.get("item_id"),item.get("slot_uk") or item.get("slot_key"),item.get("name_uk") or item.get("name"),item.get("rarity_uk") or item.get("rarity"),item.get("item_level"),item.get("refine_level"),item.get("durability"),st.get("str",0),st.get("dex",0),st.get("int",0),st.get("wis",0),st.get("sta",0),st.get("pdef",0),st.get("mdef",0),st.get("weapon_min_damage",0),st.get("weapon_max_damage",0),_runes(item),len(item.get("runes",[]) or [])])
 _table(eq,"FinalEquipmentTable");_widths(eq)
 loot=wb.create_sheet("Loot Lifecycle");_setup(loot,"Loot Lifecycle");lh=["item_id","name","category","slot","item_level","armor_type","rarity","found","equipped","sold","used_or_transformed","remaining","sale_gold"];loot.append(lh)
 for x in payload.get("loot_lifecycle",[]) or []:loot.append([x.get("item_id"),x.get("name_uk") or x.get("name_raw"),x.get("category"),x.get("slot_uk") or x.get("slot_key"),x.get("item_level"),x.get("armor_type_uk") or x.get("armor_type"),x.get("rarity_uk") or x.get("rarity"),x.get("found"),x.get("equipped"),x.get("sold"),x.get("used_or_transformed"),x.get("remaining"),x.get("sale_gold")])
 _table(loot,"LootLifecycleTable");_widths(loot)
 eco=wb.create_sheet("Economy");_setup(eco,"Economy");eco.append(["Показник","Сума"]);eco_rows=[("Стартове золото",0.0),("Золото з мобів",economy.get("from_mobs",0)),("Золото з продажу",economy.get("from_item_sales",0)),("Витрати на ремонт",-float(economy.get("repair_cost",0) or 0)),("Розрахункове фінальне золото",f"=SUM(B2:B5)"),("Фактичне фінальне золото",economy.get("final",0)),("Різниця",f"=B6-B7")]
 for row in eco_rows:eco.append(row)
 for c in eco["B"]:c.number_format='#,##0.00;[Red](#,##0.00);-'
 _table(eco,"EconomyTable");_widths(eco)
 wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.calculation.calcMode="auto";wb.save(output);return output

def _final_detail(canonical,key):
 for e in reversed(canonical.get("events",[]) or []):
  if e.get("event_type")=="simulation_final_state":return (e.get("details") or {}).get(key,0)
 return 0
